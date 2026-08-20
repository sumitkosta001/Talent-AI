"""python-docx based DOCX Resume Text Extractor.

Extracts plain text and document elements, preserving paragraph boundaries,
heading styles, table structures, and document metadata.
"""

import io
import logging
import re
from typing import List
import docx

from app.exceptions.resume import ResumeParsingError
from .models import (
    ExtractedDocument,
    DocumentParagraph,
    DocumentTable,
    DocumentTableRow,
    DocumentTableCell,
    DocumentMetadata,
)

logger = logging.getLogger("talentai.resume_processing.docx")

MIN_EXTRACTABLE_CHAR_THRESHOLD = 15


def _is_heading_style(style_name: str | None) -> tuple[bool, int | None]:
    """Check if paragraph style denotes a heading and return heading level."""
    if not style_name:
        return False, None
    name_clean = style_name.strip().lower()
    if "title" in name_clean:
        return True, 0
    if "subtitle" in name_clean:
        return True, 1
    match = re.search(r"heading\s*(\d+)", name_clean)
    if match:
        return True, int(match.group(1))
    return False, None


def extract_docx_text(content: bytes) -> ExtractedDocument:
    """Extract text, paragraphs, headings, tables, and metadata from a DOCX binary stream.

    Args:
        content: Raw binary bytes of the DOCX document.

    Returns:
        ExtractedDocument containing full text, paragraphs, tables, and metadata.

    Raises:
        ResumeParsingError: If document is empty, corrupted, or cannot be parsed as a DOCX.
    """
    if not content or len(content.strip()) == 0:
        logger.warning("DOCX extraction failed: file content is empty.")
        raise ResumeParsingError("The uploaded DOCX file is empty.")

    try:
        stream = io.BytesIO(content)
        doc = docx.Document(stream)
    except Exception as exc:
        logger.warning("python-docx failed to open document: %s", str(exc))
        raise ResumeParsingError("Failed to open or parse corrupted DOCX document.") from exc

    try:
        paragraphs: List[DocumentParagraph] = []
        full_text_parts: List[str] = []
        total_chars = 0
        total_words = 0

        # Extract paragraphs in order
        for p_idx, p in enumerate(doc.paragraphs):
            p_text = p.text.strip()
            style_name = p.style.name if p.style and hasattr(p.style, "name") else None
            is_heading, heading_level = _is_heading_style(style_name)

            if p_text:
                paragraphs.append(
                    DocumentParagraph(
                        text=p_text,
                        paragraph_index=p_idx,
                        style_name=style_name,
                        is_heading=is_heading,
                        heading_level=heading_level,
                    )
                )
                full_text_parts.append(p_text)
                total_chars += len(p_text)
                total_words += len(p_text.split())

        # Extract tables
        tables: List[DocumentTable] = []
        for t_idx, table in enumerate(doc.tables):
            table_rows: List[DocumentTableRow] = []
            table_text_lines: List[str] = []

            for r_idx, row in enumerate(table.rows):
                row_cells: List[DocumentTableCell] = []
                row_texts: List[str] = []

                for c_idx, cell in enumerate(row.cells):
                    c_text = cell.text.strip()
                    row_cells.append(
                        DocumentTableCell(
                            text=c_text,
                            row_index=r_idx,
                            col_index=c_idx,
                        )
                    )
                    if c_text:
                        row_texts.append(c_text)

                table_rows.append(DocumentTableRow(row_index=r_idx, cells=row_cells))
                if row_texts:
                    table_text_lines.append(" | ".join(row_texts))

            table_text_block = "\n".join(table_text_lines)
            if table_text_block:
                tables.append(
                    DocumentTable(
                        table_index=t_idx,
                        rows=table_rows,
                        text=table_text_block,
                    )
                )
                full_text_parts.append(table_text_block)
                total_chars += len(table_text_block)
                total_words += len(table_text_block.split())

        # Extract core metadata properties
        meta = DocumentMetadata()
        try:
            core_props = doc.core_properties
            if core_props:
                meta = DocumentMetadata(
                    title=core_props.title or None,
                    author=core_props.author or None,
                    subject=core_props.subject or None,
                    keywords=core_props.keywords or None,
                    creator=core_props.last_modified_by or None,
                    creation_date=str(core_props.created) if core_props.created else None,
                    modification_date=str(core_props.modified) if core_props.modified else None,
                )
        except Exception as e:
            logger.debug("Could not read DOCX core properties: %s", str(e))

        full_text = "\n\n".join(full_text_parts) if full_text_parts else ""
        has_extractable_text = total_chars >= MIN_EXTRACTABLE_CHAR_THRESHOLD

        logger.info(
            "DOCX extraction completed: paragraphs=%d, tables=%d, characters=%d, words=%d, has_text=%s",
            len(paragraphs), len(tables), total_chars, total_words, has_extractable_text,
        )

        return ExtractedDocument(
            text=full_text,
            document_type="docx",
            extraction_method="python-docx",
            page_count=1,
            paragraphs=paragraphs,
            tables=tables,
            metadata=meta,
            character_count=total_chars,
            word_count=total_words,
            has_extractable_text=has_extractable_text,
        )

    except Exception as exc:
        if isinstance(exc, ResumeParsingError):
            raise
        logger.warning("Unexpected error during DOCX parsing: %s", str(exc))
        raise ResumeParsingError("Failed to extract content from DOCX document.") from exc
